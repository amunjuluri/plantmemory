"""Import maintenance history exported from a CMMS (SAP PM IW29/IW39 lists, Maximo, Excel).

Gemini looks at the header and a few rows and says which column holds which field. Code
then applies that mapping to every row, so no record is paraphrased or invented. Rows
come back for review and are saved only when confirmed (same as logbook photos).
"""

import csv
import io
import json
import re
from datetime import date, datetime

from app.agent import prompts
from app.agent.tools import AgentDeps
from app.intake.machine_logs import decode
from app.models import CmmsImportResult, CmmsMapping, IncidentPublic, Industry, Outcome

CMMS_ID_PREFIX = "CMMS"
MAX_ROWS = 500
MAX_BYTES = 5 * 1024 * 1024
_DATE_FORMATS = ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d.%m.%y", "%Y%m%d")
_RECURRED = re.compile(r"\b(again|repeat(ed)?|recurr(ed|ence)|came back|re-?occur)", re.I)


class InvalidExportError(ValueError):
    """The upload is not a readable CSV or Excel export."""


def read_rows(data: bytes, filename: str | None) -> tuple[list[str], list[dict[str, str]]]:
    """Header and rows from a CSV, semicolon list, tab-separated text or .xlsx file."""
    if len(data) > MAX_BYTES:
        raise InvalidExportError("Export is larger than 5 MB. Export a shorter date range.")
    if data.startswith(b"PK\x03\x04"):
        return _read_xlsx(data)
    text = decode(data)
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if len(lines) < 2:
        raise InvalidExportError("The export needs a header row and at least one data row.")
    try:
        dialect = csv.Sniffer().sniff("\n".join(lines[:20]), delimiters=";,\t|")
    except csv.Error as err:
        raise InvalidExportError("Couldn't read that file as a table. Export as CSV.") from err
    reader = csv.reader(io.StringIO("\n".join(lines)), dialect)
    header, *body = list(reader)
    return _rows(header, body)


def _read_xlsx(data: bytes) -> tuple[list[str], list[dict[str, str]]]:
    from openpyxl import load_workbook

    try:
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
    except Exception as err:
        raise InvalidExportError("Couldn't open that Excel file.") from err
    values = [
        ["" if v is None else _cell(v) for v in row]
        for row in sheet.iter_rows(values_only=True)
        if any(v not in (None, "") for v in row)
    ]
    if len(values) < 2:
        raise InvalidExportError("The sheet needs a header row and at least one data row.")
    return _rows(values[0], values[1:])


def _cell(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _rows(header: list[str], body: list[list[str]]) -> tuple[list[str], list[dict[str, str]]]:
    names = [h.strip() or f"Column {i + 1}" for i, h in enumerate(header)]
    rows = [
        {names[i]: (row[i].strip() if i < len(row) else "") for i in range(len(names))}
        for row in body[:MAX_ROWS]
    ]
    return names, rows


def propose_mapping(header: list[str], rows: list[dict[str, str]], deps: AgentDeps) -> CmmsMapping:
    """Ask the model which column holds which field, then drop any column it made up."""
    prompt = prompts.CMMS_PROMPT.format(
        columns=json.dumps(header, ensure_ascii=False),
        rows=json.dumps(rows[:4], ensure_ascii=False, indent=1),
    )
    mapping = deps.llm.generate(prompt, CmmsMapping, thinking="low")
    known = set(header)

    def keep(cols: list[str]) -> list[str]:
        return [c for c in cols if c in known]

    def one(col: str | None) -> str | None:
        return col if col in known else None

    if mapping.date_column not in known:
        raise InvalidExportError("Couldn't find a date column in that export.")
    return mapping.model_copy(
        update={
            "id_column": one(mapping.id_column),
            "equipment_columns": keep(mapping.equipment_columns),
            "description_columns": keep(mapping.description_columns),
            "cause_columns": keep(mapping.cause_columns),
            "action_columns": keep(mapping.action_columns),
            "downtime_column": one(mapping.downtime_column),
            "status_column": one(mapping.status_column),
        }
    )


def parse_date(text: str) -> date | None:
    text = text.strip().split(" ")[0]
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_minutes(text: str, unit: str) -> int | None:
    """'2,5' hours -> 150; '1:30' -> 90; '45' minutes -> 45."""
    text = text.strip()
    if not text:
        return None
    clock = re.fullmatch(r"(\d+):(\d{2})(?::\d{2})?", text)
    if clock:
        return int(clock.group(1)) * 60 + int(clock.group(2))
    number = re.search(r"\d+(?:[.,]\d+)?", text)
    if not number:
        return None
    value = float(number.group(0).replace(",", "."))
    factor = {"minutes": 1, "hours": 60, "days": 1440}[unit]
    return max(0, round(value * factor))


def _join(row: dict[str, str], cols: list[str], sep: str = " / ") -> str:
    seen: list[str] = []
    for c in cols:
        value = row.get(c, "").strip()
        if sep == ". ":
            value = value.rstrip(". ")
        if value and value not in seen:
            seen.append(value)
    return sep.join(seen)


def _record_id(row: dict[str, str], mapping: CmmsMapping, index: int) -> str:
    raw = row.get(mapping.id_column or "", "").strip()
    token = re.sub(r"[^A-Za-z0-9]", "", raw)[:20] or f"ROW{index + 1}"
    return f"{CMMS_ID_PREFIX}-{token.upper()}"


def apply_mapping(
    rows: list[dict[str, str]],
    mapping: CmmsMapping,
    source_ref: str,
    industry: Industry | None,
) -> tuple[list[IncidentPublic], list[str]]:
    """Turn every row into a record for review, exactly as written in the export."""
    records: list[IncidentPublic] = []
    skipped: list[str] = []
    done = {v.strip().upper() for v in mapping.completed_values}
    taken: set[str] = set()
    for i, row in enumerate(rows):
        when = parse_date(row.get(mapping.date_column, ""))
        description = _join(row, mapping.description_columns, ". ")
        if when is None:
            skipped.append(f"Row {i + 2}: no readable date.")
            continue
        if not description:
            skipped.append(f"Row {i + 2}: no description.")
            continue
        status = row.get(mapping.status_column or "", "").upper()
        outcome: Outcome = (
            "recurred"
            if _RECURRED.search(description)
            else "resolved"
            if done and any(v in status.split() or v == status for v in done)
            else "unknown"
        )
        record_id = _record_id(row, mapping, i)
        if record_id in taken:
            record_id = f"{record_id}-{i + 1}"
        taken.add(record_id)
        downtime = (
            parse_minutes(row.get(mapping.downtime_column, ""), mapping.downtime_unit)
            if mapping.downtime_column
            else None
        )
        records.append(
            IncidentPublic(
                id=record_id,
                industry=industry or "pharma",
                date=when,
                area_or_equipment=_join(row, mapping.equipment_columns) or "unspecified",
                description=description,
                root_cause=_join(row, mapping.cause_columns, ". ") or None,
                action_taken=_join(row, mapping.action_columns, ". ") or None,
                outcome=outcome,
                downtime_minutes=downtime,
                source="cmms_import",
                source_ref=source_ref,
                confidence_flags=["outcome"] if outcome == "unknown" else [],
            )
        )
    return records, skipped


def import_export(
    data: bytes, filename: str | None, deps: AgentDeps, industry: Industry | None = None
) -> CmmsImportResult:
    """Read a CMMS export into records for review. Nothing is saved yet."""
    header, rows = read_rows(data, filename)
    if not rows:
        raise InvalidExportError("The export has no data rows.")
    mapping = propose_mapping(header, rows, deps)
    source_ref = re.sub(r"[^\w.\- ]", "_", filename or "cmms-export.csv")[:120]
    records, skipped = apply_mapping(rows, mapping, source_ref, industry)
    system = mapping.system or "maintenance system export"
    return CmmsImportResult(
        source_ref=source_ref,
        records=records,
        page_notes=f"Read as {system}: {len(records)} of {len(rows)} rows ready to check.",
        mapping=mapping,
        rows=len(rows),
        skipped=skipped,
    )
